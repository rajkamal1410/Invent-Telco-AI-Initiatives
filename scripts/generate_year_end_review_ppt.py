from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "year-end-review" / "Year_End_Review_2026.pptx"
SOURCE_URL = (
    "https://capgemini-my.sharepoint.com/personal/rajkamal_b_a_capgemini_com/"
    "_layouts/15/Doc.aspx?sourcedoc=%7B5BB61CDE-17FC-432E-991A-426092A55989%7D"
    "&file=YER.docx&action=default&mobileredirect=true"
)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(31, 40, 61)
MUTED = RGBColor(100, 108, 126)
INDIGO = RGBColor(70, 79, 235)
BLUE = RGBColor(0, 112, 173)
CYAN = RGBColor(18, 171, 219)
PALE = RGBColor(246, 247, 252)
LAVENDER = RGBColor(239, 240, 255)
LINE = RGBColor(230, 230, 230)
GREEN = RGBColor(30, 142, 106)


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP,
            margin=0, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        kind, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.8)
    return shape


def base_slide(title=None, kicker=None, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    rect(slide, 0, 0, 13.333, 0.06, BLUE)
    rect(slide, 0, 0, 0.72, 0.06, CYAN)
    if title:
        if kicker:
            textbox(slide, 0.7, 0.42, 11.8, 0.24, kicker.upper(), 9, INDIGO, True)
        textbox(slide, 0.7, 0.77, 12, 0.55, title, 27, INK, True)
    rect(slide, 0.7, 7.07, 11.93, 0.012, LINE)
    textbox(slide, 0.7, 7.15, 5, 0.18, "YEAR-END REVIEW  |  2026", 8, MUTED, True)
    textbox(slide, 11.8, 7.13, 0.8, 0.2, f"{number:02d}" if number else "", 9, INDIGO, True, PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title, body=None, accent=INDIGO,
         fill=WHITE, title_size=14, body_size=11):
    rect(slide, x, y, w, h, fill, LINE, True)
    rect(slide, x, y, 0.055, h, accent)
    textbox(slide, x + 0.22, y + 0.18, w - 0.42, 0.34, title, title_size, INK, True)
    if body:
        textbox(slide, x + 0.22, y + 0.62, w - 0.42, h - 0.76, body, body_size, MUTED)


def metric(slide, x, y, w, value, label, detail, accent=INDIGO):
    rect(slide, x, y, w, 1.54, WHITE, LINE, True)
    textbox(slide, x + 0.2, y + 0.16, w - 0.4, 0.58, value, 30, accent, True)
    textbox(slide, x + 0.2, y + 0.78, w - 0.4, 0.27, label, 11, INK, True)
    textbox(slide, x + 0.2, y + 1.12, w - 0.4, 0.22, detail, 9, MUTED)


def bullet_list(slide, x, y, w, items, size=12, color=MUTED, gap=0.46):
    for index, item in enumerate(items):
        yy = y + index * gap
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.09), Inches(0.08), Inches(0.08)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = INDIGO
        dot.line.fill.background()
        textbox(slide, x + 0.2, yy, w - 0.2, gap - 0.04, item, size, color)


def slide_cover():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK
    rect(slide, 0, 0, 13.333, 0.075, BLUE)
    rect(slide, 0, 0, 0.85, 0.075, CYAN)
    rect(slide, 9.7, 0.075, 3.633, 7.425, RGBColor(37, 45, 80))
    rect(slide, 10.1, 0.55, 2.55, 0.055, CYAN)
    textbox(slide, 0.9, 0.85, 7.9, 0.3, "CAPGEMINI INVENT  •  YEAR-END REVIEW", 11, CYAN, True)
    textbox(slide, 0.9, 1.75, 8.3, 1.45, "Building impact\nbeyond delivery", 37, WHITE, True)
    textbox(
        slide, 0.9, 3.55, 7.6, 0.95,
        "Client value. AI-led innovation. Reusable capability.",
        19, RGBColor(220, 226, 242)
    )
    rect(slide, 0.9, 5.18, 5.55, 0.012, RGBColor(82, 91, 126))
    textbox(slide, 0.9, 5.48, 6.3, 0.38, "2026 CONTRIBUTION SUMMARY", 12, WHITE, True)
    textbox(slide, 0.9, 5.93, 6.7, 0.3, "Account Owner  |  Cox account", 13, RGBColor(190, 199, 222))
    textbox(slide, 10.13, 1.05, 2.4, 0.3, "YEAR IN REVIEW", 10, RGBColor(190, 199, 222), True)
    textbox(slide, 10.08, 1.58, 2.5, 0.7, "93%", 37, WHITE, True)
    textbox(slide, 10.12, 2.35, 2.45, 0.54, "revenue target\nattainment", 13, RGBColor(220, 226, 242))
    rect(slide, 10.12, 3.32, 2.5, 0.012, RGBColor(82, 91, 126))
    textbox(slide, 10.08, 3.72, 2.5, 0.7, "€200K", 31, WHITE, True)
    textbox(slide, 10.12, 4.42, 2.45, 0.56, "qualified growth\npipeline", 13, RGBColor(220, 226, 242))
    textbox(slide, 10.12, 6.7, 2.5, 0.2, "2026", 10, CYAN, True)


def slide_executive_snapshot():
    slide = base_slide("A year of impact across four dimensions", "Executive snapshot", 2)
    metrics = [
        ("€214K", "REVENUE DELIVERED", "93% of €230K target", INDIGO),
        ("~€200K", "QUALIFIED PIPELINE", "Built through account mining", BLUE),
        ("100%", "ARVE UTILIZATION", "Sustained through the year", GREEN),
        ("6", "STRATEGIC ROLE AREAS", "Demand identified across the account", CYAN),
    ]
    for i, values in enumerate(metrics):
        metric(slide, 0.7 + i * 3.02, 1.58, 2.78, *values)
    pillars = [
        ("01", "Client value", "Protected account stability and trusted relationships amid merger-driven constraints."),
        ("02", "AI innovation", "Created governance visibility and a centralized knowledge platform concept."),
        ("03", "Practice building", "Produced reusable agile delivery and knowledge-management assets."),
        ("04", "People & brand", "Shared AI capability and strengthened telecom thought leadership."),
    ]
    for i, (n, title, body) in enumerate(pillars):
        x = 0.7 + (i % 2) * 6.03
        y = 3.55 + (i // 2) * 1.35
        rect(slide, x, y, 5.7, 1.12, PALE, LINE, True)
        textbox(slide, x + 0.2, y + 0.2, 0.55, 0.28, n, 12, INDIGO, True)
        textbox(slide, x + 0.9, y + 0.17, 4.5, 0.26, title, 14, INK, True)
        textbox(slide, x + 0.9, y + 0.5, 4.45, 0.49, body, 10, MUTED)


def slide_account():
    slide = base_slide("Steady account leadership through market headwinds", "01  |  Client & business", 3)
    textbox(slide, 0.72, 1.52, 5.8, 0.28, "THE CONTEXT", 10, INDIGO, True)
    textbox(
        slide, 0.72, 1.9, 5.55, 0.9,
        "Cox–Charter merger activity, tighter discretionary spend, delayed decisions and rate pressure challenged growth.",
        17, INK, True
    )
    bullet_list(
        slide, 0.75, 3.15, 5.45,
        [
            "Maintained client engagement and trusted-advisor visibility.",
            "Managed resources proactively and onboarded a new DFRC resource.",
            "Identified demand across six strategic role areas.",
        ],
        12, MUTED, 0.65
    )
    rect(slide, 6.72, 1.55, 5.88, 4.82, PALE, LINE, True)
    textbox(slide, 7.05, 1.88, 5.1, 0.32, "REVENUE PERFORMANCE", 10, INDIGO, True)
    textbox(slide, 7.05, 2.34, 3.1, 0.65, "€214K", 34, INK, True)
    textbox(slide, 9.4, 2.52, 2.6, 0.3, "of €230K target", 13, MUTED)
    rect(slide, 7.05, 3.2, 4.95, 0.22, LINE, LINE, True)
    rect(slide, 7.05, 3.2, 4.6, 0.22, INDIGO, INDIGO, True)
    textbox(slide, 7.05, 3.56, 1.0, 0.25, "93%", 12, INDIGO, True)
    textbox(slide, 7.05, 4.2, 4.8, 0.3, "GROWTH PIPELINE", 10, INDIGO, True)
    textbox(slide, 7.05, 4.57, 2.2, 0.55, "~€200K", 27, INK, True)
    textbox(slide, 9.3, 4.66, 2.65, 0.53, "Qualified opportunities through account mining and leadership engagement", 10, MUTED)
    textbox(slide, 7.05, 5.55, 4.95, 0.56,
            "Cloud Platform • Enterprise Architecture • Power Platform • DevOps • ProdOps • Delivery Leadership",
            10, MUTED)


def slide_ai():
    slide = base_slide("AI governance turned fragmented activity into shared visibility", "02  |  AI-led innovation", 4)
    textbox(slide, 0.72, 1.5, 11.7, 0.35,
            "Proactively designed an AI Governance Dashboard and Knowledge Management Platform for the SFA domain.",
            15, INK, True)
    cards = [
        ("GOVERNANCE DASHBOARD", "Centralized visibility into AI investments, adoption, benefits realization and initiative progress.", INDIGO),
        ("AI KNOWLEDGE LIBRARY", "A single source of truth for TPMs, architects, delivery teams and business stakeholders.", BLUE),
        ("STANDARDIZED REPORTING", "Common documentation and reporting across domains supports clearer accountability.", CYAN),
    ]
    for i, (title, body, accent) in enumerate(cards):
        card(slide, 0.72 + i * 4.02, 2.25, 3.72, 2.0, title, body, accent, PALE, 12, 12)
    rect(slide, 0.72, 4.75, 11.76, 1.35, LAVENDER, LAVENDER, True)
    textbox(slide, 1.0, 4.98, 2.2, 0.25, "HOW IT WAS BUILT", 10, INDIGO, True)
    textbox(slide, 3.3, 4.94, 8.55, 0.7,
            "GitHub Copilot • AI-assisted development • Vibe Coding practices",
            17, INK, True, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.98, 5.58, 10.9, 0.28,
            "Value: stronger executive oversight, easier knowledge discovery and a reusable governance approach.",
            11, MUTED)


def slide_recognition():
    slide = base_slide("Client trust and telecom thought leadership", "03  |  Recognition & brand", 5)
    rect(slide, 0.72, 1.55, 5.55, 4.75, INK, INK, True)
    textbox(slide, 1.05, 1.92, 4.85, 0.27, "CLIENT RECOGNITION", 10, CYAN, True)
    textbox(slide, 1.05, 2.42, 4.65, 1.42,
            "Recognized by client leadership for AI adoption, stakeholder alignment, delivery excellence and dependable program support.",
            20, WHITE, True)
    textbox(slide, 1.05, 4.45, 4.65, 0.75,
            "A signal of trusted partnership across multiple strategic initiatives.",
            12, RGBColor(211, 218, 234))
    textbox(slide, 6.85, 1.65, 5.3, 0.3, "THOUGHT LEADERSHIP", 10, INDIGO, True)
    textbox(slide, 6.85, 2.12, 5.5, 1.12,
            "“Intelligent Telco: Monetizing Data in the AI Era”",
            23, INK, True)
    bullet_list(
        slide, 6.88, 3.55, 5.35,
        [
            "Published an external perspective on data monetization and AI-driven insight.",
            "Connected telecom data assets to new revenue and digital ecosystem opportunities.",
            "Created reusable IP for account mining, proposals and executive discussions.",
            "Raised Capgemini Invent visibility in telecom transformation and AI.",
        ],
        11, MUTED, 0.59
    )


def slide_assets():
    slide = base_slide("Reusable assets make delivery more consistent and scalable", "04  |  Practice building", 6)
    card(
        slide, 0.72, 1.55, 5.76, 4.8, "AGILE DELIVERY WAYS OF WORKING TOOLKIT",
        "A reusable framework for delivery engagements, mobilization and proposal responses.",
        INDIGO, PALE, 14, 12
    )
    tags = [
        ("Roadmaps", 0.98, 2.82, 1.42), ("Iteration planning", 2.52, 2.82, 1.95),
        ("Backlog refinement", 0.98, 3.38, 2.08), ("Health checks", 3.18, 3.38, 1.48),
        ("Status tracking", 0.98, 3.94, 1.65), ("Retrospectives", 2.76, 3.94, 1.63),
    ]
    for label, x, y, w in tags:
        rect(slide, x, y, w, 0.38, WHITE, LINE, True)
        textbox(slide, x + 0.08, y + 0.08, w - 0.16, 0.2, label, 9, INDIGO, True, PP_ALIGN.CENTER)
    textbox(slide, 0.98, 4.75, 5.0, 0.9,
            "Designed to improve governance consistency and reduce setup effort through ready-to-use templates.",
            11, MUTED)
    card(
        slide, 6.82, 1.55, 5.78, 4.8, "INVENT TMT KNOWLEDGE REPOSITORY",
        "LUMOS-inspired concept to make organizational IP easier to find and reuse.",
        BLUE, WHITE, 14, 12
    )
    textbox(slide, 7.08, 2.82, 4.95, 0.3, "CENTRALIZED CONTENT", 10, INDIGO, True)
    bullet_list(
        slide, 7.1, 3.23, 4.95,
        ["Proposal responses & business development assets",
         "Points of view & delivery accelerators",
         "Consulting collateral & reusable knowledge"],
        11, MUTED, 0.57
    )
    textbox(slide, 7.08, 5.18, 4.92, 0.76,
            "Designed to support faster onboarding, proposal preparation and client engagement readiness.",
            11, MUTED)


def slide_team():
    slide = base_slide("Capability building extended impact beyond individual delivery", "05  |  Team impact", 7)
    metric(slide, 0.72, 1.6, 2.8, "100%", "ARVE UTILIZATION", "Sustained across the year", GREEN)
    textbox(slide, 4.0, 1.65, 7.8, 0.45, "AI capability shared through practical enablement", 18, INK, True)
    bullet_list(
        slide, 4.03, 2.35, 8.0,
        [
            "Promoted AI-assisted delivery and GitHub Copilot usage.",
            "Shared guidance on prompt engineering and responsible AI use.",
            "Facilitated discussions on AI-native operating models and context engineering.",
            "Explored prompt structures and AI-versus-technical-debt trade-offs.",
        ],
        12, MUTED, 0.68
    )
    rect(slide, 0.72, 3.65, 2.8, 2.55, LAVENDER, LAVENDER, True)
    textbox(slide, 0.98, 3.97, 2.25, 0.3, "BALANCED CONTRIBUTION", 10, INDIGO, True)
    textbox(slide, 0.98, 4.45, 2.24, 1.2,
            "Client delivery\n+ innovation\n+ team enablement",
            16, INK, True)
    textbox(slide, 4.03, 5.68, 7.8, 0.38,
            "Outcome: stronger awareness of AI-enabled productivity and practical adoption across consulting and delivery.",
            11, MUTED)


def slide_close():
    slide = base_slide("Impact that compounds into the next year", "06  |  Summary & next horizon", 8)
    textbox(slide, 0.72, 1.52, 11.7, 0.72,
            "Client value • business growth • AI-led innovation • practice development",
            21, INDIGO, True)
    columns = [
        ("PROTECT", "Sustain trusted client relationships as merger-related change continues."),
        ("CONVERT", "Progress the ~€200K opportunity pipeline into qualified demand and revenue."),
        ("SCALE", "Extend governance, knowledge and delivery assets for broader reuse."),
        ("MEASURE", "Track adoption and realized outcomes from AI and enablement initiatives."),
    ]
    for i, (title, body) in enumerate(columns):
        x = 0.72 + i * 3.02
        rect(slide, x, 2.75, 2.78, 1.95, PALE, LINE, True)
        textbox(slide, x + 0.2, 3.02, 2.35, 0.3, title, 12, INDIGO, True)
        textbox(slide, x + 0.2, 3.48, 2.35, 0.94, body, 11, MUTED)
    rect(slide, 0.72, 5.12, 11.84, 1.25, INK, INK, True)
    textbox(slide, 1.02, 5.39, 10.95, 0.56,
            "A year of sustained leadership—delivering measurable account outcomes while building capabilities that can be reused.",
            16, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    source = textbox(slide, 8.58, 6.6, 3.95, 0.23, "Source: supplied YER summary", 8, INDIGO, True, PP_ALIGN.RIGHT)
    source.text_frame.paragraphs[0].runs[0].hyperlink.address = SOURCE_URL


def main():
    slide_cover()
    slide_executive_snapshot()
    slide_account()
    slide_ai()
    slide_recognition()
    slide_assets()
    slide_team()
    slide_close()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
