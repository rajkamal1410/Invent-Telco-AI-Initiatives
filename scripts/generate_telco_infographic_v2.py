from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUT = r"c:\Users\rajkamaa\Downloads\AI exercises\deliverables\intelligent-telco\The Intelligent Telco - Enhanced Infographic.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(7, 19, 31)
NAVY_2 = RGBColor(19, 33, 52)
PANEL = RGBColor(18, 35, 52)
TEXT = RGBColor(236, 244, 255)
MUTED = RGBColor(184, 203, 223)
CYAN = RGBColor(34, 211, 238)
BLUE = RGBColor(96, 165, 250)
PURPLE = RGBColor(139, 92, 246)
GREEN = RGBColor(52, 211, 153)
AMBER = RGBColor(251, 191, 36)
RED = RGBColor(248, 113, 113)


def set_bg(slide, c1=NAVY, c2=NAVY_2):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = c1
    glow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    glow.fill.solid(); glow.fill.fore_color.rgb = c1; glow.line.fill.background()
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.55))
    band.fill.solid(); band.fill.fore_color.rgb = c2; band.line.fill.background()


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=TEXT,
                align=PP_ALIGN.LEFT, font_name='Aptos', italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = 0
    run = p.runs[0] if p.runs else p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_card(slide, left, top, width, height, fill=PANEL, line=RGBColor(148,163,184), radius=0.12):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid(); card.fill.fore_color.rgb = fill
    card.line.color.rgb = line
    card.line.width = Pt(1)
    return card


def add_circle(slide, cx, cy, r, fill, line=RGBColor(148,163,184), line_width=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx-r, cy-r, r*2, r*2)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_arrow(slide, left, top, width, height, color=CYAN):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arrow.fill.solid(); arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()
    return arrow


def add_bullets(slide, left, top, width, height, items, size=12, color=MUTED):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.name = 'Aptos'


def add_title(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(8), Inches(0.5), title, size=24, bold=True)
    if subtitle:
        add_textbox(slide, Inches(0.8), Inches(1.0), Inches(9), Inches(0.4), subtitle, size=11, color=CYAN, bold=True)


def slide1():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    add_card(slide, Inches(0.7), Inches(0.4), Inches(12.0), Inches(0.8), fill=RGBColor(27, 46, 63), line=RGBColor(96,165,250))
    add_textbox(slide, Inches(1.1), Inches(0.58), Inches(0.7), Inches(0.4), 'IT', size=26, bold=True, color=RGBColor(7,19,31), align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.9), Inches(0.58), Inches(7.5), Inches(0.4), 'INTELLIGENT TELCO', size=24, bold=True)

    add_textbox(slide, Inches(0.9), Inches(1.55), Inches(5.0), Inches(0.5), 'DATA MONETIZATION IN THE AI ERA', size=12, bold=True, color=CYAN)
    add_textbox(slide, Inches(0.9), Inches(2.0), Inches(11.8), Inches(1.8), 'The', size=42, bold=True)
    add_textbox(slide, Inches(0.9), Inches(3.0), Inches(12.0), Inches(1.6), 'INTELLIGENT TELCO', size=54, bold=True, color=RGBColor(103, 179, 255))
    add_textbox(slide, Inches(0.9), Inches(4.8), Inches(11.4), Inches(1.6),
                'Operators are no longer just connectivity providers. They sit on trusted network signals, identity data, mobility patterns and enterprise-grade intelligence that can be transformed into monetizable products, programmable APIs and decision-ready services.',
                size=21, color=TEXT)

    k1 = add_card(slide, Inches(0.9), Inches(6.1), Inches(5.7), Inches(0.8), fill=RGBColor(24, 38, 52), line=RGBColor(148, 163, 184))
    add_textbox(slide, Inches(1.25), Inches(6.35), Inches(2.0), Inches(0.3), 'Connectivity', size=26, bold=True)
    add_textbox(slide, Inches(1.25), Inches(6.7), Inches(2.0), Inches(0.2), 'LEGACY CORE', size=9, color=MUTED)

    k2 = add_card(slide, Inches(6.9), Inches(6.1), Inches(5.7), Inches(0.8), fill=RGBColor(24, 38, 52), line=RGBColor(148,163,184))
    add_textbox(slide, Inches(7.2), Inches(6.35), Inches(2.0), Inches(0.3), 'Data', size=26, bold=True)
    add_textbox(slide, Inches(7.2), Inches(6.7), Inches(2.0), Inches(0.2), 'RAW SIGNAL', size=9, color=MUTED)


def slide2():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 'Why AI changes the economics', 'Operators already own the assets; AI turns them into monetizable outcomes')

    # central narrative circle
    add_circle(slide, Inches(6.7), Inches(3.7), Inches(1.8), fill=RGBColor(17, 31, 48), line=CYAN, line_width=2)
    add_textbox(slide, Inches(5.65), Inches(2.95), Inches(2.1), Inches(0.5), 'Network', size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(5.65), Inches(3.45), Inches(2.1), Inches(0.5), 'Signals', size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(5.65), Inches(4.05), Inches(2.1), Inches(0.5), 'Identity', size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(7.75), Inches(2.95), Inches(2.1), Inches(0.5), 'Mobility', size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(7.75), Inches(3.45), Inches(2.1), Inches(0.5), 'Quality', size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(7.75), Inches(4.05), Inches(2.1), Inches(0.5), 'Trust', size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

    add_card(slide, Inches(0.8), Inches(2.0), Inches(2.8), Inches(1.3), fill=RGBColor(16, 31, 48))
    add_textbox(slide, Inches(1.0), Inches(2.2), Inches(2.2), Inches(0.4), 'Asset base', size=15, bold=True)
    add_textbox(slide, Inches(1.0), Inches(2.6), Inches(2.2), Inches(0.6), 'Device data\nSubscriber relationships\nNetwork quality\nEnterprise access', size=11, color=MUTED)

    add_card(slide, Inches(0.8), Inches(4.2), Inches(2.8), Inches(1.3), fill=RGBColor(16, 31, 48))
    add_textbox(slide, Inches(1.0), Inches(4.4), Inches(2.2), Inches(0.4), 'AI enablement', size=15, bold=True)
    add_textbox(slide, Inches(1.0), Inches(4.8), Inches(2.2), Inches(0.6), 'Forecasting\nOptimization\nFraud detection\nPersonalization', size=11, color=MUTED)

    add_card(slide, Inches(9.7), Inches(2.0), Inches(2.8), Inches(1.3), fill=RGBColor(16, 31, 48))
    add_textbox(slide, Inches(9.9), Inches(2.2), Inches(2.2), Inches(0.4), 'Commercial value', size=15, bold=True)
    add_textbox(slide, Inches(9.9), Inches(2.6), Inches(2.2), Inches(0.6), 'Revenue uplift\nLower churn\nLower fraud costs\nNew API products', size=11, color=MUTED)

    add_card(slide, Inches(9.7), Inches(4.2), Inches(2.8), Inches(1.3), fill=RGBColor(16, 31, 48))
    add_textbox(slide, Inches(9.9), Inches(4.4), Inches(2.2), Inches(0.4), 'Outcome shift', size=15, bold=True)
    add_textbox(slide, Inches(9.9), Inches(4.8), Inches(2.2), Inches(0.6), 'From data seller\nto intelligence provider', size=11, color=MUTED)


def slide3():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 'Two monetization paths', 'Internal efficiency and external productization')

    left = Inches(0.8)
    widths = [Inches(5.7), Inches(5.7)]
    y1 = Inches(1.8)
    y2 = Inches(4.6)

    add_card(slide, left, y1, widths[0], Inches(2.1), fill=RGBColor(19, 36, 53))
    add_textbox(slide, left + Inches(0.3), y1 + Inches(0.2), Inches(2.8), Inches(0.4), '1. Internal monetization', size=20, bold=True, color=CYAN)
    add_bullets(slide, left + Inches(0.3), y1 + Inches(0.7), Inches(5.0), Inches(1.1), [
        'Improve customer personalization and next-best actions',
        'Reduce cost-to-serve via AI-driven network optimization',
        'Increase retention and lifetime value with better service outcomes',
    ], size=12)

    add_card(slide, left + Inches(6.5), y1, widths[1], Inches(2.1), fill=RGBColor(19, 36, 53))
    add_textbox(slide, left + Inches(6.8), y1 + Inches(0.2), Inches(2.8), Inches(0.4), '2. External monetization', size=20, bold=True, color=GREEN)
    add_bullets(slide, left + Inches(6.8), y1 + Inches(0.7), Inches(5.0), Inches(1.1), [
        'Network APIs and programmable services',
        'Identity, authentication and fraud intelligence',
        'Enterprise analysis and mobility intelligence',
    ], size=12)

    add_card(slide, Inches(2.6), Inches(4.5), Inches(8.2), Inches(1.6), fill=RGBColor(20, 42, 63))
    add_textbox(slide, Inches(3.0), Inches(4.9), Inches(7.2), Inches(0.5), 'Bottom line:', size=16, bold=True, color=AMBER)
    add_textbox(slide, Inches(3.0), Inches(5.35), Inches(7.0), Inches(0.5), 'The most scalable telco model is not selling raw data—it is productizing trusted intelligence and network capabilities.', size=18, color=TEXT)


def slide4():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 'Network APIs as the programmable platform', 'From connectivity product to business capability layer')

    # left stack
    add_card(slide, Inches(0.8), Inches(1.8), Inches(3.3), Inches(4.3), fill=RGBColor(19, 36, 53))
    add_textbox(slide, Inches(1.1), Inches(2.1), Inches(2.5), Inches(0.4), 'API portfolio', size=18, bold=True)
    add_bullets(slide, Inches(1.05), Inches(2.6), Inches(2.7), Inches(2.8), [
        'Number Verification',
        'SIM Swap',
        'Device Swap',
        'KYC / identity checks',
        'Location verification',
        'Quality on Demand',
        'Device status',
        'Edge-related capabilities',
    ], size=11)

    # center / example stack
    add_card(slide, Inches(4.6), Inches(1.8), Inches(3.7), Inches(4.3), fill=RGBColor(20, 41, 59))
    add_textbox(slide, Inches(4.9), Inches(2.1), Inches(2.5), Inches(0.4), 'Why it matters', size=18, bold=True, color=CYAN)
    add_bullets(slide, Inches(4.9), Inches(2.6), Inches(3.1), Inches(2.8), [
        'Reframes network as a programmable platform',
        'Turns internal capabilities into external products',
        'Creates reusable developer and enterprise services',
        'Improves economics beyond connectivity alone',
        'Supports authentication, fraud prevention and trust use cases',
    ], size=11)

    # right case study
    add_card(slide, Inches(8.8), Inches(1.8), Inches(3.8), Inches(4.3), fill=RGBColor(17, 31, 48))
    add_textbox(slide, Inches(9.1), Inches(2.1), Inches(2.7), Inches(0.4), 'Example: US market', size=18, bold=True, color=GREEN)
    add_textbox(slide, Inches(9.1), Inches(2.7), Inches(3.1), Inches(1.2), 'AT&T, T-Mobile and Verizon have standardized Number Verification and SIM Swap APIs via Aduna to support identity and fraud workflows.', size=12, color=TEXT)
    add_textbox(slide, Inches(9.1), Inches(4.0), Inches(2.9), Inches(0.5), 'Commercial impact', size=13, bold=True, color=AMBER)
    add_textbox(slide, Inches(9.1), Inches(4.5), Inches(3.0), Inches(0.9), 'New revenue streams from authentication, fraud prevention and trusted service execution.', size=12, color=MUTED)


def slide5():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 'Identity, fraud and enterprise intelligence', 'Trusted signals drive defensible telco value')

    add_card(slide, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.0), fill=RGBColor(19, 36, 53))
    add_textbox(slide, Inches(1.1), Inches(2.1), Inches(3.5), Inches(0.4), 'Identity and fraud intelligence', size=20, bold=True, color=CYAN)
    add_bullets(slide, Inches(1.1), Inches(2.7), Inches(5.0), Inches(2.2), [
        'Operators hold trusted relationships with devices and subscriptions',
        'Number Verification can confirm device-to-number association',
        'SIM Swap capabilities indicate recent changes and risk events',
        'This creates strong value for fintech, digital commerce and e-commerce trust workflows',
    ], size=12)

    add_card(slide, Inches(6.9), Inches(1.8), Inches(5.7), Inches(4.0), fill=RGBColor(19, 36, 53))
    add_textbox(slide, Inches(7.2), Inches(2.1), Inches(3.8), Inches(0.4), 'Enterprise mobility intelligence', size=20, bold=True, color=GREEN)
    add_bullets(slide, Inches(7.2), Inches(2.7), Inches(5.0), Inches(2.2), [
        'Aggregated mobility and network intelligence support planning decisions',
        'Use cases include retail, logistics, transport, public sector and branch strategy',
        'Vodafone and Telefónica illustrate how anonymized data can become decision support products',
        'The value is in better enterprise decisions, not raw customer data resale',
    ], size=12)


def slide6():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 'Governance and strategic playbook', 'Trust must be designed into the product itself')

    add_card(slide, Inches(0.8), Inches(1.8), Inches(3.5), Inches(4.2), fill=RGBColor(19, 36, 53))
    add_textbox(slide, Inches(1.1), Inches(2.1), Inches(2.6), Inches(0.4), 'Governance must include', size=16, bold=True)
    add_bullets(slide, Inches(1.1), Inches(2.7), Inches(3.0), Inches(2.7), [
        'Data quality and provenance',
        'Privacy, purpose and consent',
        'Model governance and monitoring',
        'Cybersecurity and API protection',
        'Explainability and traceability',
        'Human accountability',
    ], size=11)

    add_card(slide, Inches(4.7), Inches(1.8), Inches(3.8), Inches(4.2), fill=RGBColor(18, 40, 60))
    add_textbox(slide, Inches(5.1), Inches(2.1), Inches(3.0), Inches(0.4), '4-step path to scale', size=18, bold=True, color=CYAN)
    add_bullets(slide, Inches(5.1), Inches(2.7), Inches(3.1), Inches(2.9), [
        'Prove value internally',
        'Embed privacy and governance',
        'Productize differentiated network signals',
        'Scale through ecosystems and APIs',
    ], size=12)

    add_card(slide, Inches(8.9), Inches(1.8), Inches(3.7), Inches(4.2), fill=RGBColor(17, 31, 48))
    add_textbox(slide, Inches(9.2), Inches(2.1), Inches(2.8), Inches(0.4), 'Strategic outcome', size=18, bold=True, color=GREEN)
    add_textbox(slide, Inches(9.2), Inches(2.7), Inches(2.6), Inches(2.4), 'Connectivity → Data → Insights → Intelligence Products → AI-enabled ecosystems', size=18, bold=True, color=TEXT)


def slide7():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_textbox(slide, Inches(0.9), Inches(0.8), Inches(11.8), Inches(0.7), 'The Intelligent Telco', size=34, bold=True, color=TEXT)
    add_textbox(slide, Inches(0.9), Inches(1.7), Inches(11.8), Inches(0.6), 'Trusted intelligence is the next growth engine for telecom', size=20, color=CYAN, bold=True)

    add_card(slide, Inches(0.8), Inches(2.7), Inches(3.8), Inches(2.3), fill=RGBColor(17, 31, 48))
    add_textbox(slide, Inches(1.1), Inches(3.2), Inches(3.0), Inches(0.4), 'Network APIs', size=17, bold=True)
    add_textbox(slide, Inches(1.1), Inches(3.7), Inches(3.0), Inches(0.9), 'Programmable, reusable capabilities that turn infrastructure into products.', size=12, color=MUTED)

    add_card(slide, Inches(4.9), Inches(2.7), Inches(3.8), Inches(2.3), fill=RGBColor(17, 31, 48))
    add_textbox(slide, Inches(5.2), Inches(3.2), Inches(3.0), Inches(0.4), 'Identity & trust', size=17, bold=True)
    add_textbox(slide, Inches(5.2), Inches(3.7), Inches(3.0), Inches(0.9), 'Authentication, fraud prevention and trusted digital services.', size=12, color=MUTED)

    add_card(slide, Inches(9.0), Inches(2.7), Inches(3.5), Inches(2.3), fill=RGBColor(17, 31, 48))
    add_textbox(slide, Inches(9.3), Inches(3.2), Inches(2.8), Inches(0.4), 'Enterprise intelligence', size=17, bold=True)
    add_textbox(slide, Inches(9.3), Inches(3.7), Inches(2.8), Inches(0.9), 'Mobility, planning and decision support products for business customers.', size=12, color=MUTED)

    add_textbox(slide, Inches(0.9), Inches(5.8), Inches(11.2), Inches(0.6), 'The winning telco strategy is to monetize trusted intelligence—not raw data.', size=18, bold=True, color=TEXT)


slide1(); slide2(); slide3(); slide4(); slide5(); slide6(); slide7()

# ensure folder exists
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print(f"Created: {OUT}")
