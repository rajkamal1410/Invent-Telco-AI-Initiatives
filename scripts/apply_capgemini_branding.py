from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "documents" / "Intelligent_Telco_Infographic_Rich_Presentation.pptx"
OUTPUT = (
    ROOT
    / "deliverables"
    / "intelligent-telco"
    / "Intelligent_Telco_Infographic_Rich_Presentation_Capgemini_Branded.pptx"
)

BRAND_BLUE = "0070AD"
BRIGHT_BLUE = "0058AB"
DEEP_BLUE = "004280"
CYAN = "12ABDB"
LIGHT_BLUE = "77BDFF"
PALE_BLUE = "EAF3FB"
CHARCOAL = "272936"
BODY = "383B4E"
MUTED = "626D7A"
NAVY = "082B4C"

FILL_COLORS = {
    "0B63B6": BRAND_BLUE,
    "0796A5": CYAN,
    "6741C4": DEEP_BLUE,
    "238A55": BRIGHT_BLUE,
    "A1508A": DEEP_BLUE,
    "A94F12": DEEP_BLUE,
    "FFF2E6": PALE_BLUE,
    "E8F8F9": PALE_BLUE,
    "F1EDF9": PALE_BLUE,
    "A8EBF0": LIGHT_BLUE,
    "66D6DF": CYAN,
    "E3F0F5": PALE_BLUE,
}
FONT_COLORS = {
    **FILL_COLORS,
    "082B4C": CHARCOAL,
    "102A43": CHARCOAL,
    "627487": MUTED,
    "D2F1F3": PALE_BLUE,
}

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
THEME_COLORS = {
    "dk1": CHARCOAL,
    "lt1": "FFFFFF",
    "dk2": BODY,
    "lt2": "F6F9FC",
    "accent1": BRAND_BLUE,
    "accent2": CYAN,
    "accent3": BRIGHT_BLUE,
    "accent4": DEEP_BLUE,
    "accent5": LIGHT_BLUE,
    "accent6": "4DC0E4",
    "hlink": BRAND_BLUE,
    "folHlink": DEEP_BLUE,
}


def mapped_rgb(rgb, mapping):
    if rgb is None:
        return None
    return mapping.get(str(rgb).upper())


def recolor_shape(shape):
    try:
        if shape.fill.type is not None:
            replacement = mapped_rgb(shape.fill.fore_color.rgb, FILL_COLORS)
            if replacement:
                shape.fill.fore_color.rgb = RGBColor.from_string(replacement)
    except (AttributeError, ValueError, TypeError):
        pass

    try:
        replacement = mapped_rgb(shape.line.color.rgb, FILL_COLORS)
        if replacement:
            shape.line.color.rgb = RGBColor.from_string(replacement)
    except (AttributeError, ValueError, TypeError):
        pass

    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    replacement = mapped_rgb(run.font.color.rgb, FONT_COLORS)
                    if replacement:
                        run.font.color.rgb = RGBColor.from_string(replacement)
                except (AttributeError, ValueError, TypeError):
                    pass

    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        try:
                            replacement = mapped_rgb(run.font.color.rgb, FONT_COLORS)
                            if replacement:
                                run.font.color.rgb = RGBColor.from_string(replacement)
                        except (AttributeError, ValueError, TypeError):
                            pass

    if shape.shape_type == 6:
        for child in shape.shapes:
            recolor_shape(child)


def add_branding(slide, dark_slide):
    blue_rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.045)
    )
    blue_rule.fill.solid()
    blue_rule.fill.fore_color.rgb = RGBColor.from_string(BRAND_BLUE)
    blue_rule.line.fill.background()

    cyan_rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.62), Inches(0.045)
    )
    cyan_rule.fill.solid()
    cyan_rule.fill.fore_color.rgb = RGBColor.from_string(CYAN)
    cyan_rule.line.fill.background()

    wordmark = slide.shapes.add_textbox(
        Inches(11.48), Inches(7.13), Inches(1.35), Inches(0.2)
    )
    text_frame = wordmark.text_frame
    text_frame.clear()
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = "Capgemini"
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("FFFFFF" if dark_slide else BRAND_BLUE)


def apply_theme_palette(pptx_path):
    nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with NamedTemporaryFile(
        suffix=".pptx", dir=pptx_path.parent, delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)

    try:
        with ZipFile(pptx_path, "r") as source_zip, ZipFile(
            temporary_path, "w", ZIP_DEFLATED
        ) as output_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "ppt/theme/theme1.xml":
                    root = etree.fromstring(data)
                    color_scheme = root.find(".//a:clrScheme", nsmap)
                    for name, color in THEME_COLORS.items():
                        element = color_scheme.find(f"a:{name}", nsmap)
                        if element is None:
                            continue
                        for child in list(element):
                            element.remove(child)
                        etree.SubElement(
                            element,
                            f"{{{NS['a']}}}srgbClr",
                            val=color,
                        )
                    data = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                output_zip.writestr(info, data)
        os.replace(temporary_path, pptx_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def main():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Source presentation not found: {SOURCE}")

    presentation = Presentation(SOURCE)
    for slide_index, slide in enumerate(presentation.slides):
        for shape in slide.shapes:
            recolor_shape(shape)
        add_branding(slide, dark_slide=slide_index in (0, len(presentation.slides) - 1))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    apply_theme_palette(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
