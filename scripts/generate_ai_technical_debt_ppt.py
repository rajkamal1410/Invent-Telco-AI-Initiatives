from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "ai-technical-debt" / "AI_Technical_Debt_Discussion.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(6, 22, 37)
NAVY_LIGHT = RGBColor(12, 33, 51)
CARD = RGBColor(17, 38, 56)
CARD_BEST = RGBColor(14, 55, 54)
TEXT = RGBColor(240, 247, 252)
MUTED = RGBColor(184, 205, 222)
CYAN = RGBColor(47, 205, 224)
BLUE = RGBColor(101, 166, 250)
GREEN = RGBColor(73, 220, 165)
AMBER = RGBColor(255, 195, 78)
LINE = RGBColor(43, 66, 83)


QUESTIONS = [
    {
        "title": "The CEO's AI moment",
        "scenario": (
            "After seeing competitors announce GenAI, NovaTel's CEO wants a visible AI success in six months "
            "and no more architecture-only discussion. Customer data is fragmented across 14 systems."
        ),
        "prompt": "What should you do?",
        "options": [
            {
                "text": "Launch a GenAI customer assistant immediately.",
                "upside": "Meets the CEO's urgency and creates visible excitement.",
                "risk": "With fragmented data, success may be superficial and hard to scale.",
            },
            {
                "text": "Refuse and fix the data foundations first.",
                "upside": "Data quality matters and can prevent future rework.",
                "risk": "Value may wait 12-18 months, weakening executive sponsorship.",
            },
            {
                "text": "Choose one journey, improve only its required data, and deploy AI.",
                "upside": "A focused journey (such as billing disputes) can prove business value.",
                "risk": "Requires tight scope and outcome measures to avoid a one-off pilot.",
            },
            {
                "text": "Build a company-wide AI platform before starting use cases.",
                "upside": "A shared platform can support enterprise scale and consistency.",
                "risk": "Risks building technology before validating business needs.",
            },
        ],
        "best": "C",
        "why": (
            "It balances urgency with sustainability: deliver a measurable outcome while fixing only the "
            "technical debt that blocks that journey."
        ),
    },
    {
        "title": "The perfect AI model nobody uses",
        "scenario": (
            "A churn model scores 88% accuracy, but Sales does not trust it, some customer data is incomplete, "
            "and it is not connected to CRM. The board wants rollout approval."
        ),
        "prompt": "What do you do?",
        "options": [
            {
                "text": "Deploy immediately; the accuracy is good enough.",
                "upside": "Puts the model into use quickly so it can create value.",
                "risk": "Untrusted recommendations and weak data can damage adoption.",
            },
            {
                "text": "Wait until data quality reaches 100%.",
                "upside": "Reduces risk from incomplete or unreliable inputs.",
                "risk": "Perfection can become a reason to delay learning and value indefinitely.",
            },
            {
                "text": "Deploy in one region and monitor the results.",
                "upside": "A controlled rollout creates evidence and learning.",
                "risk": "Keep the scope bounded and monitor trust, outcomes, and data issues.",
            },
            {
                "text": "Start a CRM modernization project first.",
                "upside": "Addresses the integration gap and underlying workflow constraints.",
                "risk": "A broad modernization can defer business value for years.",
            },
        ],
        "best": "C",
        "why": (
            "A regional rollout limits exposure while testing real-world performance, data quality, and user "
            "trust before deciding how to scale."
        ),
    },
    {
        "title": "Budget shock",
        "scenario": (
            "The CFO has approved ₹100 crore against a ₹300 crore request. You must decide how to invest the "
            "smaller budget."
        ),
        "prompt": "How should you allocate it?",
        "options": [
            {
                "text": "Spend it all on data modernization.",
                "upside": "Builds a stronger foundation for future initiatives.",
                "risk": "May produce no visible business wins in the near term.",
            },
            {
                "text": "Spend it all on AI initiatives.",
                "upside": "Signals fast innovation and focuses on new capabilities.",
                "risk": "AI may amplify existing debt and fragile foundations.",
            },
            {
                "text": "Split the budget equally between data and AI.",
                "upside": "Feels balanced and avoids choosing one side.",
                "risk": "Equal funding ignores differences in use-case value and need.",
            },
            {
                "text": "Fund high-value AI use cases and the debt they specifically need fixed.",
                "upside": "Connects investment in AI and foundations to business outcomes.",
                "risk": "Needs disciplined prioritization and benefits tracking.",
            },
        ],
        "best": "D",
        "why": (
            "Investment follows business outcomes: prioritize the most valuable use cases and repair only the "
            "technical debt needed to deliver them."
        ),
    },
    {
        "title": "The legacy system dilemma",
        "scenario": (
            "Customer provisioning runs on a 15-year-old system with 99.98% reliability and good data. An "
            "architect says it must be replaced before the company can become AI-native."
        ),
        "prompt": "What should be done?",
        "options": [
            {
                "text": "Replace the system immediately.",
                "upside": "Modernization may improve flexibility over the long term.",
                "risk": "Age alone does not prove debt; replacement carries cost and delivery risk.",
            },
            {
                "text": "Leave the system untouched forever.",
                "upside": "Preserves a reliable service and avoids unnecessary disruption.",
                "risk": "It could become a constraint as integration needs change.",
            },
            {
                "text": "Wrap it with APIs and use AI where needed.",
                "upside": "Enables integration while preserving a stable, reliable core.",
                "risk": "Reassess if the system actually blocks outcomes or becomes costly to maintain.",
            },
            {
                "text": "Move everything to cloud before using AI.",
                "upside": "Cloud can support strategic modernization goals.",
                "risk": "Large investment and migration risk arrive before a business outcome.",
            },
        ],
        "best": "C",
        "why": (
            "A system is not technical debt just because it is old. Add an API boundary and modernize further "
            "only when the system constrains business outcomes."
        ),
    },
    {
        "title": "You are the new CTO",
        "scenario": (
            "The board asks whether the company should become AI-first or data-first. You can make only one "
            "statement."
        ),
        "prompt": "What is your answer?",
        "options": [
            {
                "text": "AI will help clean up data over time.",
                "upside": "AI can assist with some data-quality tasks.",
                "risk": "AI cannot magically create trusted data or reliable sources of truth.",
            },
            {
                "text": "Complete the data foundation before any AI begins.",
                "upside": "Strong foundations matter for dependable AI.",
                "risk": "A sequential approach delays learning and the market will not wait.",
            },
            {
                "text": "Run AI and data programs independently.",
                "upside": "Separate workstreams can provide clear ownership.",
                "risk": "Independent programs risk silos and misaligned priorities.",
            },
            {
                "text": "Advance AI and data together, driven by business use cases.",
                "upside": "Combines foundation-building with delivery and learning.",
                "risk": "Requires shared ownership and prioritization around outcomes.",
            },
        ],
        "best": "D",
        "why": (
            "Neither AI-first nor foundation-first: use-case-first. Progress AI and data together, with business "
            "outcomes guiding both."
        ),
    },
]


def set_background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.11)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = CYAN
    band.line.fill.background()


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=16,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_footer(slide, page, label="AI TECHNICAL DEBT  /  DISCUSSION"):
    add_text(
        slide,
        Inches(0.55),
        Inches(7.16),
        Inches(10.5),
        Inches(0.18),
        label,
        size=8,
        color=MUTED,
    )
    add_text(
        slide,
        Inches(12.05),
        Inches(7.12),
        Inches(0.7),
        Inches(0.24),
        f"{page:02d}",
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_header(slide, eyebrow, title, page):
    add_text(
        slide,
        Inches(0.62),
        Inches(0.36),
        Inches(11.8),
        Inches(0.28),
        eyebrow.upper(),
        size=10,
        color=CYAN,
        bold=True,
    )
    add_text(
        slide,
        Inches(0.62),
        Inches(0.68),
        Inches(12.1),
        Inches(0.57),
        title,
        size=27,
        bold=True,
    )
    add_footer(slide, page)


def add_cover():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)

    add_text(
        slide,
        Inches(0.86),
        Inches(0.85),
        Inches(5.7),
        Inches(0.34),
        "FACILITATED LEADERSHIP DISCUSSION",
        size=11,
        color=CYAN,
        bold=True,
    )
    add_text(
        slide,
        Inches(0.82),
        Inches(1.6),
        Inches(11.5),
        Inches(1.65),
        "AI ambition.\nTechnical debt. Real outcomes.",
        size=36,
        bold=True,
    )
    add_text(
        slide,
        Inches(0.88),
        Inches(3.62),
        Inches(9.8),
        Inches(0.92),
        "Five leadership scenarios for choosing where AI, data, and modernization create value.",
        size=20,
        color=MUTED,
    )

    accent = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.88),
        Inches(5.28),
        Inches(11.55),
        Inches(0.92),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = NAVY_LIGHT
    accent.line.color.rgb = LINE
    add_text(
        slide,
        Inches(1.15),
        Inches(5.47),
        Inches(11.0),
        Inches(0.48),
        "CORE IDEA   Use cases first. Build the data and technology foundation as outcomes demand.",
        size=16,
        color=TEXT,
        bold=True,
    )
    add_footer(slide, 1)


def add_facilitation_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Facilitator guide", "Make the answer a discussion, not a quiz", 2)

    steps = [
        ("01", "POLL", "Give participants 30 seconds to choose an option."),
        ("02", "HEAR THE MAJORITY", "Ask: “Why did you choose this?”"),
        ("03", "HEAR THE MINORITY", "Ask: “What risk do you see that others are missing?”"),
        ("04", "REVEAL", "Explore why each option is reasonable, where it falls short, and the recommended answer."),
    ]
    top = 1.65
    for number, title, description in steps:
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.82),
            Inches(top),
            Inches(11.7),
            Inches(0.88),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = CARD
        box.line.color.rgb = LINE
        add_text(
            slide,
            Inches(1.08),
            Inches(top + 0.16),
            Inches(0.7),
            Inches(0.44),
            number,
            size=19,
            color=CYAN,
            bold=True,
        )
        add_text(
            slide,
            Inches(2.0),
            Inches(top + 0.14),
            Inches(2.35),
            Inches(0.28),
            title,
            size=12,
            color=TEXT,
            bold=True,
        )
        add_text(
            slide,
            Inches(4.18),
            Inches(top + 0.13),
            Inches(7.95),
            Inches(0.56),
            description,
            size=13,
            color=MUTED,
        )
        top += 1.12

    add_text(
        slide,
        Inches(0.92),
        Inches(6.24),
        Inches(11.6),
        Inches(0.48),
        "Use the recommendation to prompt debate; the context and evidence still matter.",
        size=14,
        color=AMBER,
        bold=True,
    )


def add_question_slide(question, index):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    page = index + 2
    add_header(slide, f"Question {index} of 5", question["title"], page)

    scenario = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.62),
        Inches(1.4),
        Inches(12.1),
        Inches(0.92),
    )
    scenario.fill.solid()
    scenario.fill.fore_color.rgb = NAVY_LIGHT
    scenario.line.color.rgb = LINE
    add_text(
        slide,
        Inches(0.86),
        Inches(1.49),
        Inches(11.6),
        Inches(0.43),
        question["scenario"],
        size=12,
        color=MUTED,
    )
    add_text(
        slide,
        Inches(0.88),
        Inches(1.94),
        Inches(11.4),
        Inches(0.26),
        question["prompt"],
        size=14,
        color=TEXT,
        bold=True,
    )
    add_text(
        slide,
        Inches(0.78),
        Inches(2.38),
        Inches(3.7),
        Inches(0.24),
        "OPTION",
        size=9,
        color=CYAN,
        bold=True,
    )
    add_text(
        slide,
        Inches(4.82),
        Inches(2.38),
        Inches(7.4),
        Inches(0.24),
        "WHY IT'S REASONABLE  /  TRADE-OFF",
        size=9,
        color=CYAN,
        bold=True,
    )

    top = 2.68
    row_height = 0.78
    gap = 0.08
    for option_index, option in enumerate(question["options"]):
        letter = chr(ord("A") + option_index)
        is_best = letter == question["best"]
        row = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.62),
            Inches(top),
            Inches(12.1),
            Inches(row_height),
        )
        row.fill.solid()
        row.fill.fore_color.rgb = CARD_BEST if is_best else CARD
        row.line.color.rgb = GREEN if is_best else LINE
        row.line.width = Pt(1.5 if is_best else 0.8)

        marker = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.78),
            Inches(top + 0.14),
            Inches(0.48),
            Inches(0.48),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = GREEN if is_best else NAVY_LIGHT
        marker.line.fill.background()
        add_text(
            slide,
            Inches(0.78),
            Inches(top + 0.14),
            Inches(0.48),
            Inches(0.48),
            letter,
            size=15,
            color=NAVY if is_best else CYAN,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            Inches(1.43),
            Inches(top + 0.08),
            Inches(3.2),
            Inches(0.62),
            option["text"],
            size=11.5,
            color=TEXT,
            bold=is_best,
        )
        rationale = f"WHY: {option['upside']}\nTRADE-OFF: {option['risk']}"
        add_text(
            slide,
            Inches(4.82),
            Inches(top + 0.06),
            Inches(7.55),
            Inches(0.67),
            rationale,
            size=10.5,
            color=TEXT if is_best else MUTED,
        )
        top += row_height + gap

    answer = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.62),
        Inches(6.22),
        Inches(12.1),
        Inches(0.7),
    )
    answer.fill.solid()
    answer.fill.fore_color.rgb = RGBColor(13, 48, 48)
    answer.line.color.rgb = GREEN
    add_text(
        slide,
        Inches(0.88),
        Inches(6.31),
        Inches(1.3),
        Inches(0.45),
        f"BEST: {question['best']}",
        size=13,
        color=GREEN,
        bold=True,
    )
    add_text(
        slide,
        Inches(2.2),
        Inches(6.29),
        Inches(10.2),
        Inches(0.48),
        question["why"],
        size=11.5,
        color=TEXT,
    )


def add_close():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "Takeaway", "Use-case first—not AI first or foundation first", 8)

    statements = [
        ("START", "Choose a business outcome worth pursuing."),
        ("FOCUS", "Fix the data and technical debt that blocks that outcome."),
        ("LEARN", "Deploy with controls, measure results, and earn trust."),
        ("SCALE", "Expand what works; modernize further when evidence warrants it."),
    ]
    top = 1.62
    for label, text in statements:
        add_text(
            slide,
            Inches(1.05),
            Inches(top),
            Inches(1.55),
            Inches(0.48),
            label,
            size=12,
            color=CYAN,
            bold=True,
        )
        add_text(
            slide,
            Inches(2.72),
            Inches(top),
            Inches(9.3),
            Inches(0.5),
            text,
            size=19,
            color=TEXT,
            bold=True,
        )
        top += 1.03

    add_text(
        slide,
        Inches(1.02),
        Inches(6.15),
        Inches(11.0),
        Inches(0.52),
        "AI and data modernization move together when business outcomes set the direction.",
        size=15,
        color=AMBER,
        bold=True,
    )


add_cover()
add_facilitation_slide()
for question_index, question in enumerate(QUESTIONS, start=1):
    add_question_slide(question, question_index)
add_close()

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Created {OUT}")
