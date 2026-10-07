from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR


OUT = r"c:\Users\rajkamaa\Downloads\AI exercises\deliverables\intelligent-telco\The Intelligent Telco Infographic Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(6, 22, 37)
NAVY_2 = RGBColor(11, 31, 49)
CARD = RGBColor(18, 35, 52)
CARD_2 = RGBColor(21, 43, 62)
TEXT = RGBColor(237, 246, 255)
MUTED = RGBColor(184, 205, 225)
CYAN = RGBColor(34, 211, 238)
BLUE = RGBColor(96, 165, 250)
PURPLE = RGBColor(139, 92, 246)
GREEN = RGBColor(52, 211, 153)
AMBER = RGBColor(251, 191, 36)


def set_bg(slide, color1=NAVY, color2=NAVY_2):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color1
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color1
    rect.line.fill.background()
    # add subtle glow band
    gl = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.9))
    gl.fill.solid()
    gl.fill.fore_color.rgb = color2
    gl.line.fill.background()


def add_textbox(slide, left, top, width, height, text, size=20, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
                font_name='Aptos', italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_card(slide, left, top, width, height, title, subtitle, accent=BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(17, 31, 48)
    card.line.color.rgb = RGBColor(148, 163, 184)
    card.line.width = Pt(1)

    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.18), top + Inches(0.18), width - Inches(0.36), Inches(0.08))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    title_box = add_textbox(slide, left + Inches(0.22), top + Inches(0.38), width - Inches(0.4), Inches(0.42), title, size=18, bold=True, color=TEXT)
    subtitle_box = add_textbox(slide, left + Inches(0.22), top + Inches(0.8), width - Inches(0.4), Inches(0.9), subtitle, size=10, color=MUTED)
    return card


def add_bullet_list(slide, left, top, width, height, items, color=MUTED):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.space_after = Pt(7)
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            run.font.name = 'Aptos'
            run.font.size = Pt(15)
            run.font.color.rgb = color


def add_value_chain(slide):
    steps = [
        ("Connectivity", "Legacy core of mobile access, messaging and secure digital services.", BLUE),
        ("Data", "Network signals, mobility patterns, subscriber behavior and service quality.", CYAN),
        ("Insights", "Aggregated intelligence that supports decisions across retail, transport and enterprises.", GREEN),
        ("Products", "APIs, fraud protection, identity services and telecom data products.", PURPLE),
        ("Ecosystems", "AI orchestration, partnerships and programmable services at scale.", AMBER),
    ]
    x_positions = [Inches(0.7), Inches(2.95), Inches(5.2), Inches(7.45), Inches(9.7)]
    for i, (title, desc, accent) in enumerate(steps):
        left = x_positions[i]
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3.0), Inches(2.0), Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(16, 31, 48)
        card.line.color.rgb = accent
        card.line.width = Pt(1.5)
        num = slide.shapes.add_textbox(left + Inches(0.2), Inches(3.18), Inches(0.28), Inches(0.22))
        tf = num.text_frame
        p = tf.paragraphs[0]
        p.text = str(i + 1)
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = accent
        title_box = add_textbox(slide, left + Inches(0.15), Inches(3.55), Inches(1.7), Inches(0.35), title, size=17, bold=True)
        desc_box = add_textbox(slide, left + Inches(0.15), Inches(3.95), Inches(1.7), Inches(0.8), desc, size=9, color=MUTED)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left + Inches(2.05), Inches(3.75), Inches(0.7), Inches(0.18))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = accent; arrow.line.fill.background()


def add_icon(slide, left, top, label, icon, accent=BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.4), Inches(1.85))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(16, 31, 48); card.line.color.rgb = RGBColor(148,163,184)
    icon_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(0.5), Inches(0.5))
    icon_tf = icon_box.text_frame
    icon_tf.paragraphs[0].text = icon
    for run in icon_tf.paragraphs[0].runs:
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = accent
    title_box = add_textbox(slide, left + Inches(0.18), top + Inches(0.74), Inches(2.1), Inches(0.35), label, size=15, bold=True)


def add_slide_1():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.45), Inches(12.3), Inches(0.9))
    top_bar.fill.solid(); top_bar.fill.fore_color.rgb = RGBColor(27, 46, 63); top_bar.line.fill.background()

    square = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(0.6), Inches(0.8), Inches(0.55))
    square.fill.solid(); square.fill.fore_color.rgb = RGBColor(96, 165, 250); square.line.fill.background()
    text = add_textbox(slide, Inches(1.9), Inches(0.62), Inches(5.4), Inches(0.5), 'IT', size=26, bold=True, color=RGBColor(6, 22, 37))
    title = add_textbox(slide, Inches(2.8), Inches(0.62), Inches(7.3), Inches(0.5), 'INTELLIGENT TELCO', size=25, bold=True)

    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.62), Inches(4.7), Inches(0.52))
    pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor(20, 41, 58); pill.line.color.rgb = RGBColor(96,165,250)
    add_textbox(slide, Inches(1.02), Inches(1.76), Inches(4.2), Inches(0.3), 'DATA MONETIZATION IN THE AI ERA', size=13, bold=True, color=CYAN)

    main = add_textbox(slide, Inches(0.85), Inches(2.15), Inches(11.8), Inches(2.3), 'The\nINTELLIGENT TELCO', size=38, bold=True, color=TEXT)

    big = main.text_frame.paragraphs[0]
    big.text = 'The'
    big.font.size = Pt(50)
    big.font.bold = True
    big.font.color.rgb = TEXT
    p2 = main.text_frame.add_paragraph()
    p2.text = 'INTELLIGENT TELCO'
    p2.alignment = PP_ALIGN.LEFT
    for run in p2.runs:
        run.font.size = Pt(62)
        run.font.bold = True
        run.font.color.rgb = RGBColor(99, 175, 255)

    add_textbox(slide, Inches(0.9), Inches(4.85), Inches(11.5), Inches(1.4),
                'Telecom operators are no longer just connectivity providers. They sit on trusted network signals, identity data, mobility patterns, and enterprise-grade intelligence that can be transformed into monetizable products, programmable APIs, and decision-ready services.',
                size=22, color=TEXT)

    c1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(5.8), Inches(0.9))
    c1.fill.solid(); c1.fill.fore_color.rgb = RGBColor(22, 38, 53); c1.line.color.rgb = RGBColor(148,163,184)
    add_textbox(slide, Inches(1.1), Inches(6.25), Inches(2.5), Inches(0.4), 'Connectivity', size=23, bold=True)
    add_textbox(slide, Inches(1.1), Inches(6.65), Inches(2.5), Inches(0.2), 'LEGACY CORE', size=10, color=MUTED)

    c2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(6.1), Inches(5.8), Inches(0.9))
    c2.fill.solid(); c2.fill.fore_color.rgb = RGBColor(22, 38, 53); c2.line.color.rgb = RGBColor(148,163,184)
    add_textbox(slide, Inches(7.2), Inches(6.25), Inches(2.5), Inches(0.4), 'Data', size=23, bold=True)
    add_textbox(slide, Inches(7.2), Inches(6.65), Inches(2.5), Inches(0.2), 'RAW SIGNAL', size=10, color=MUTED)


def add_slide_2():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(6.5), Inches(0.4), 'VALUE PATHWAY', size=12, bold=True, color=CYAN)
    add_textbox(slide, Inches(0.8), Inches(0.9), Inches(8.0), Inches(0.6), 'From connectivity to intelligence', size=28, bold=True)
    add_textbox(slide, Inches(8.1), Inches(1.05), Inches(4.3), Inches(0.7), 'The winning direction is not selling raw data; it is packaging trusted signals into repeatable business outcomes.', size=11, color=MUTED)
    add_value_chain(slide)


def add_slide_3():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_textbox(slide, Inches(0.8), Inches(0.6), Inches(6.5), Inches(0.4), 'PRIMARY MONETIZATION LEVERS', size=12, bold=True, color=CYAN)
    add_textbox(slide, Inches(0.8), Inches(1.0), Inches(8.0), Inches(0.6), 'Where value is created', size=28, bold=True)

    add_icon(slide, Inches(0.8), Inches(1.9), 'Network APIs', '📡', BLUE)
    add_icon(slide, Inches(3.6), Inches(1.9), 'Identity & Fraud', '🔐', CYAN)
    add_icon(slide, Inches(6.4), Inches(1.9), 'Mobility Insights', '📍', GREEN)
    add_icon(slide, Inches(9.2), Inches(1.9), 'AI Ecosystem', '🤖', PURPLE)

    desc_y = Inches(3.9)
    for i, text_block in enumerate([
        'Turn internal network capabilities into programmable services through GSMA Open Gateway and ecosystem APIs.',
        'Use number verification, SIM swap and device trust to create premium identity and fraud prevention services.',
        'Package aggregated movement and network insights into retail, logistics and public sector decision support products.',
        'Combine telecom data with AI agents and broader digital experiences to scale engagement across ecosystems.'
    ]):
        x_positions = [Inches(0.8), Inches(3.6), Inches(6.4), Inches(9.2)]
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_positions[i], desc_y, Inches(2.35), Inches(1.45))
        panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(17, 31, 48); panel.line.color.rgb = RGBColor(148,163,184)
        add_textbox(slide, x_positions[i] + Inches(0.16), desc_y + Inches(0.15), Inches(2.05), Inches(1.0), text_block, size=10, color=MUTED)

    # second row of panels for metrics / proof
    panel1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.0), Inches(2.9), Inches(0.8))
    panel1.fill.solid(); panel1.fill.fore_color.rgb = RGBColor(15, 29, 45); panel1.line.color.rgb = RGBColor(96,165,250)
    add_textbox(slide, Inches(1.08), Inches(6.18), Inches(2.2), Inches(0.4), 'Network APIs', size=16, bold=True)
    add_textbox(slide, Inches(1.08), Inches(6.52), Inches(2.2), Inches(0.2), 'Programmable products', size=9, color=MUTED)

    panel2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.05), Inches(6.0), Inches(2.9), Inches(0.8))
    panel2.fill.solid(); panel2.fill.fore_color.rgb = RGBColor(15, 29, 45); panel2.line.color.rgb = RGBColor(34,211,238)
    add_textbox(slide, Inches(4.33), Inches(6.18), Inches(2.2), Inches(0.4), 'Identity & trust', size=16, bold=True)
    add_textbox(slide, Inches(4.33), Inches(6.52), Inches(2.2), Inches(0.2), 'Fraud-safe monetization', size=9, color=MUTED)

    panel3 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.3), Inches(6.0), Inches(2.9), Inches(0.8))
    panel3.fill.solid(); panel3.fill.fore_color.rgb = RGBColor(15, 29, 45); panel3.line.color.rgb = RGBColor(52,211,153)
    add_textbox(slide, Inches(7.58), Inches(6.18), Inches(2.2), Inches(0.4), 'Enterprise data', size=16, bold=True)
    add_textbox(slide, Inches(7.58), Inches(6.52), Inches(2.2), Inches(0.2), 'Mobility intelligence', size=9, color=MUTED)

    panel4 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.55), Inches(6.0), Inches(2.0), Inches(0.8))
    panel4.fill.solid(); panel4.fill.fore_color.rgb = RGBColor(15, 29, 45); panel4.line.color.rgb = RGBColor(139,92,246)
    add_textbox(slide, Inches(10.82), Inches(6.18), Inches(1.5), Inches(0.4), 'AI + ecosystem', size=15, bold=True)


def add_slide_4():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_textbox(slide, Inches(0.8), Inches(0.7), Inches(6.5), Inches(0.4), 'OPERATING MODEL', size=12, bold=True, color=CYAN)
    add_textbox(slide, Inches(0.8), Inches(1.1), Inches(8.2), Inches(0.5), 'Simple, scalable and governable', size=28, bold=True)

    steps = [
        ('01', 'Prove value internally', 'Start with customer, network and operational use cases that improve revenue, cost and experience.'),
        ('02', 'Build trust into architecture', 'Embed consent, privacy, data quality and governance from day one.'),
        ('03', 'Productize differentiated signals', 'Focus on unique network-originated capabilities that are hard to replicate.'),
        ('04', 'Scale through ecosystems', 'Use GSMA Open Gateway, Aduna and partnerships to distribute capabilities across developers and enterprises.')
    ]
    y = Inches(1.9)
    for i, (num, title, desc) in enumerate(steps):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.8), Inches(1.1))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(17,31,48); box.line.color.rgb = RGBColor(148,163,184)
        # index chip
        chip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), y + Inches(0.18), Inches(0.55), Inches(0.5))
        chip.fill.solid(); chip.fill.fore_color.rgb = RGBColor(96,165,250); chip.line.fill.background()
        add_textbox(slide, Inches(1.1), y + Inches(0.2), Inches(0.3), Inches(0.4), num, size=15, bold=True, color=RGBColor(6,22,37))
        add_textbox(slide, Inches(1.8), y + Inches(0.15), Inches(4.0), Inches(0.34), title, size=16, bold=True)
        add_textbox(slide, Inches(1.8), y + Inches(0.52), Inches(9.8), Inches(0.4), desc, size=11, color=MUTED)
        y += Inches(1.3)


def add_slide_5():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_textbox(slide, Inches(0.8), Inches(0.7), Inches(7.2), Inches(0.4), 'BOTTOM LINE', size=12, bold=True, color=CYAN)
    add_textbox(slide, Inches(0.8), Inches(1.05), Inches(10.0), Inches(0.5), 'The strategic opportunity', size=28, bold=True)

    cards = [
        ('Network as platform', 'Operators can convert network capability into programmable business products instead of static utility services.', BLUE),
        ('Trust as advantage', 'Identity, verification and fraud signals become premium assets when governed, privacy-safe and commercially clear.', CYAN),
        ('AI as enabler', 'AI agents can orchestrate services across ecosystems, but only within explainable, auditable and governed workflows.', GREEN),
    ]
    x = [Inches(0.8), Inches(4.8), Inches(8.8)]
    for i, (title, desc, accent) in enumerate(cards):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x[i], Inches(2.1), Inches(3.5), Inches(3.3))
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(17,31,48); card.line.color.rgb = accent
        card.line.width = Pt(1.5)
        add_textbox(slide, x[i] + Inches(0.35), Inches(2.48), Inches(2.8), Inches(0.36), title, size=20, bold=True)
        add_textbox(slide, x[i] + Inches(0.35), Inches(2.98), Inches(2.8), Inches(1.7), desc, size=14, color=MUTED)

    quote = add_textbox(slide, Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.5), '“The real product is not AI itself. It is trusted intelligence, capability, or business outcome that AI helps the operator create.”', size=18, bold=True, color=TEXT)
    add_textbox(slide, Inches(0.8), Inches(6.45), Inches(5.0), Inches(0.25), 'TELCO POV • STRATEGIC NARRATIVE', size=11, color=MUTED)


add_slide_1()
add_slide_2()
add_slide_3()
add_slide_4()
add_slide_5()

prs.save(OUT)
print(f"Created: {OUT}")
